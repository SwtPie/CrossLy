"""Participant import from CSV / Excel files.

Pipeline:  read_table (file -> headers + rows)  ->  detect_mapping (header -> field)
           ->  preview / commit (rows -> validated participants, duplicates flagged).

Everything that interprets a value (names, sex, dates, age, VMA, column
headers) is a small pure function, so it can be checked by hand. Only
`preview` and `commit` touch the database.
"""

import csv
import io
import re
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

from . import db, messages, participants
from .events import existing_event_path
from .normalize import collapse_spaces

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_ROWS = 5000
MAX_COLUMNS = 60
HEADER_SCAN_ROWS = 15
MAX_BLANK_RUN = 1000  # stop reading a sheet after this many empty rows in a row

FIELDS = (
    "last_name", "first_name", "full_name", "sex", "class_name",
    "establishment", "birth_date", "age", "vma",
)

# Header names (normalized: no accents, lowercase, punctuation as single spaces).
SYNONYMS: dict[str, set[str]] = {
    "last_name": {
        "nom", "nom de famille", "nom famille", "last name", "lastname", "family name",
        "surname", "nom eleve", "nom du participant", "nom participant",
    },
    "first_name": {
        "prenom", "prenoms", "first name", "firstname", "given name", "prenom eleve",
        "prenom du participant", "prenom participant",
    },
    "full_name": {
        "nom prenom", "prenom nom", "nom et prenom", "prenom et nom", "nom complet",
        "full name", "fullname", "eleve", "eleves", "participant", "identite", "name",
    },
    "sex": {"sexe", "genre", "sex", "gender", "m f", "f m", "g f", "sexe m f"},
    "class_name": {"classe", "class", "division", "cl", "classe eleve"},
    "establishment": {
        "etablissement", "etab", "ecole", "college", "lycee", "school",
        "etablissement scolaire", "structure",
    },
    "birth_date": {
        "date de naissance", "date naissance", "naissance", "ne le", "nee le",
        "birth date", "birthdate", "date of birth", "dob", "ddn", "date nais",
    },
    "age": {"age", "ages"},
    "vma": {
        "vma", "vitesse maximale aerobie", "vitesse max aerobie", "vma km h", "vma kmh",
    },
}
_EXACT = {name: field for field, names in SYNONYMS.items() for name in names}
# Weaker rule for headers that merely contain a keyword ("VMA (km/h)", "Date de naissance de l'élève").
_KEYWORDS = (
    ("vma", "vma"), ("prenom", "first_name"), ("prenoms", "first_name"),
    ("naissance", "birth_date"), ("etablissement", "establishment"),
    ("classe", "class_name"), ("sexe", "sex"), ("age", "age"), ("nom", "last_name"),
)

_EMPTY_MARKERS = {"", "-", "--", "?", "/", "nc", "na", "n a", "n/a"}
_SEX_MALE = {"m", "g", "h", "garcon", "garcons", "homme", "masculin", "male", "boy", "gars", "1"}
_SEX_FEMALE = {"f", "fille", "filles", "femme", "feminin", "female", "girl", "2"}
_DATE_FORMATS = (
    "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%Y-%m-%d", "%Y/%m/%d",
    "%d/%m/%y", "%d-%m-%y", "%d.%m.%y",
)


# --- pure helpers ----------------------------------------------------------------

def normalize_header(value: Any) -> str:
    """'Né(e) le' -> 'ne le', 'VMA (km/h)' -> 'vma km h'."""
    text = re.sub(r"\(e\)", "", str(value or ""), flags=re.IGNORECASE)
    return re.sub(r"[^a-z0-9]+", " ", db.normalize_key(text)).strip()


def match_field(normalized_header: str) -> str | None:
    if normalized_header in _EXACT:
        return _EXACT[normalized_header]
    tokens = set(normalized_header.split())
    for keyword, field in _KEYWORDS:
        if keyword in tokens:
            return field
    return None


def detect_mapping(headers: list[str]) -> dict[str, int | None]:
    """Map each field to the index of its column (first match wins), or None."""
    mapping: dict[str, int | None] = {field: None for field in FIELDS}
    for index, header in enumerate(headers):
        field = match_field(normalize_header(header))
        if field and mapping[field] is None:
            mapping[field] = index
    if mapping["last_name"] is not None and mapping["first_name"] is not None:
        mapping["full_name"] = None  # separate columns are more reliable
    return mapping


def default_options(headers: list[str]) -> dict:
    """Options suggested by the headers (a 'Prénom Nom' column means first name first)."""
    first_name_first = any(normalize_header(h) in ("prenom nom", "prenom et nom") for h in headers)
    return {"establishment": "", "class_name": "", "first_name_first": first_name_first}


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _is_empty(value: Any) -> bool:
    """Blank, or a placeholder such as '-' or 'NC'."""
    if _is_blank(value):
        return True
    return isinstance(value, str) and db.normalize_key(value) in _EMPTY_MARKERS


def cell_text(value: Any) -> str:
    """Display text of a cell: whole floats lose their '.0', dates become ISO."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    return collapse_spaces(str(value))


def smart_title(text: str) -> str:
    """Fix names exported all in capitals (or all in lowercase); mixed case is left alone."""
    text = collapse_spaces(text)
    if sum(c.isalpha() for c in text) < 2:
        return text
    if text == text.upper() or text == text.lower():
        return text.lower().title()
    return text


def split_full_name(text: str, first_name_first: bool = False) -> tuple[str, str]:
    """Split one 'NOM Prénom' (or 'Prénom NOM') cell into (last name, first name).

    Capitalized words mark the last name when the cell mixes cases; otherwise the
    first word is the last name (or the last word, with `first_name_first`).
    A comma ('Dupont, Élise') always separates last name from first name.
    """
    text = collapse_spaces(text)
    if "," in text:
        last, _, first = text.partition(",")
        return collapse_spaces(last), collapse_spaces(first)
    tokens = text.split(" ")
    if len(tokens) < 2:
        return text, ""
    if first_name_first:
        trailing = 0
        while trailing < len(tokens) and tokens[-1 - trailing].isupper():
            trailing += 1
        cut = len(tokens) - trailing if 0 < trailing < len(tokens) else len(tokens) - 1
        return " ".join(tokens[cut:]), " ".join(tokens[:cut])
    leading = 0
    while leading < len(tokens) and tokens[leading].isupper():
        leading += 1
    cut = leading if 0 < leading < len(tokens) else 1
    return " ".join(tokens[:cut]), " ".join(tokens[cut:])


def parse_sex(value: Any) -> tuple[str | None, bool]:
    """(code, ok): 'M' / 'F'. Empty gives (None, True); unknown gives (None, False)."""
    if _is_empty(value):
        return None, True
    key = re.sub(r"[^a-z0-9]+", "", db.normalize_key(cell_text(value)))
    if key in _SEX_MALE:
        return "M", True
    if key in _SEX_FEMALE:
        return "F", True
    return None, False


def parse_birth_date(value: Any) -> tuple[str | None, bool]:
    """(ISO date, ok). Accepts dd/mm/yyyy, dd-mm-yyyy, dd.mm.yyyy, yyyy-mm-dd, 2-digit years,
    real Excel dates and Excel serial numbers."""
    if _is_empty(value):
        return None, True
    parsed: date | None = None
    if isinstance(value, datetime):
        parsed = value.date()
    elif isinstance(value, date):
        parsed = value
    elif isinstance(value, (int, float)) and not isinstance(value, bool):
        if 1000 < value < 80000:  # Excel serial day number
            parsed = date(1899, 12, 30) + timedelta(days=int(value))
    else:
        text = collapse_spaces(str(value)).split(" ")[0]
        for fmt in _DATE_FORMATS:
            try:
                parsed = datetime.strptime(text, fmt).date()
                break
            except ValueError:
                continue
    if parsed is None or parsed.year < 1900 or parsed > date.today():
        return None, False
    return parsed.isoformat(), True


def parse_age(value: Any) -> tuple[int | None, bool]:
    if _is_empty(value):
        return None, True
    number: float | None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        number = value
    else:
        match = re.fullmatch(r"(\d{1,2})(?:[.,]0+)?\s*(?:ans?)?", collapse_spaces(str(value)).lower())
        number = int(match.group(1)) if match else None
    if number is None or number != int(number) or not 3 <= number <= 99:
        return None, False
    return int(number), True


def parse_vma(value: Any) -> tuple[float | None, bool]:
    """(km/h, ok). Accepts '14,5', '14.5' and '14,5 km/h'; the range 5-30 is enforced."""
    if _is_empty(value):
        return None, True
    number: float | None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        number = float(value)
    else:
        match = re.search(r"\d+(?:[.,]\d+)?", str(value))
        number = float(match.group().replace(",", ".")) if match else None
    if number is None or not 5 <= number <= 30:
        return None, False
    return round(number, 2), True


# --- reading files -------------------------------------------------------------------

def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def _detect_delimiter(text: str) -> str:
    sample = [line for line in text.splitlines() if line.strip()][:5]
    counts = {d: sum(line.count(d) for line in sample) for d in (";", ",", "\t")}
    best = max(counts, key=counts.get)
    return best if counts[best] > 0 else ";"


def _read_csv(path: Path) -> list[tuple[int, list]]:
    text = _decode(path.read_bytes())
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=_detect_delimiter(text))
    rows = []
    try:
        for cells in reader:
            cells = cells[:MAX_COLUMNS]
            if not all(_is_blank(c) for c in cells):
                rows.append((reader.line_num, cells))
    except csv.Error:
        raise ValueError(messages.IMPORT_FILE_UNREADABLE) from None
    return rows


def _read_xlsx(path: Path) -> dict[str, list[tuple[int, list]]]:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 - corrupt, encrypted or not really an .xlsx
        raise ValueError(messages.IMPORT_FILE_UNREADABLE) from None
    try:
        sheets: dict[str, list[tuple[int, list]]] = {}
        for sheet in workbook.worksheets:
            rows = []
            blank_run = 0
            for line, values in enumerate(sheet.iter_rows(values_only=True), start=1):
                cells = list(values[:MAX_COLUMNS])
                if all(_is_blank(c) for c in cells):
                    blank_run += 1
                    if blank_run > MAX_BLANK_RUN:
                        break
                    continue
                blank_run = 0
                rows.append((line, cells))
            sheets[sheet.title] = rows
        return sheets
    except Exception:  # noqa: BLE001
        raise ValueError(messages.IMPORT_FILE_UNREADABLE) from None
    finally:
        workbook.close()


def _header_score(cells: list) -> int:
    fields = {match_field(normalize_header(c)) for c in cells if isinstance(c, str)}
    return len(fields - {None})


def read_table(source: str) -> dict:
    """Read a .csv or .xlsx file into {file_name, sheet, headers, header_line, rows}.

    The header row is the one (among the first rows of the best sheet) that looks
    most like column titles, so title lines above the table are skipped. Rows are
    (line number in the file, cell values).
    """
    path = Path(source)
    suffix = path.suffix.lower()
    if suffix not in (".csv", ".xlsx", ".xlsm"):
        raise ValueError(messages.IMPORT_FILE_TYPE)
    if not path.is_file():
        raise ValueError(messages.IMPORT_FILE_UNREADABLE)
    if path.stat().st_size > MAX_FILE_BYTES:
        raise ValueError(messages.IMPORT_FILE_TOO_BIG)
    try:
        sheets = {"": _read_csv(path)} if suffix == ".csv" else _read_xlsx(path)
    except OSError:
        raise ValueError(messages.IMPORT_FILE_UNREADABLE) from None

    best = None  # (name, rows, header index, score)
    for name, rows in sheets.items():
        if not rows:
            continue
        index, score = 0, -1
        for i, (_, cells) in enumerate(rows[:HEADER_SCAN_ROWS]):
            if _header_score(cells) > score:
                index, score = i, _header_score(cells)
        if best is None or score > best[3]:
            best = (name, rows, index, score)
    if best is None:
        raise ValueError(messages.IMPORT_FILE_EMPTY)

    name, rows, index, _ = best
    header_line, header_cells = rows[index]
    data = rows[index + 1:]
    if len(data) > MAX_ROWS:
        raise ValueError(messages.IMPORT_TOO_MANY_ROWS)
    width = min(MAX_COLUMNS, max(len(cells) for _, cells in rows[index:]))
    headers = [cell_text(c) for c in header_cells] + [""] * (width - len(header_cells))
    return {
        "file_name": path.name,
        "sheet": name,
        "headers": headers,
        "header_line": header_line,
        "rows": data,
    }


# --- planning an import ---------------------------------------------------------------

def _clean_mapping(mapping: dict | None, column_count: int) -> dict[str, int | None]:
    result: dict[str, int | None] = {field: None for field in FIELDS}
    for field in FIELDS:
        value = (mapping or {}).get(field)
        if value in (None, ""):
            continue
        try:
            index = int(value)
        except (TypeError, ValueError):
            continue
        if 0 <= index < column_count:
            result[field] = index
    return result


def _clean_overrides(overrides: dict | None) -> dict[int, str]:
    result = {}
    for key, value in (overrides or {}).items():
        try:
            result[int(key)] = collapse_spaces("" if value is None else str(value))
        except (TypeError, ValueError):
            continue
    return result


def _plan(table: dict, mapping: dict | None, options: dict | None, overrides: dict | None,
          existing: set[tuple[str, str]]) -> tuple[dict, list[dict], list[str]]:
    """Validate every row. Returns (clean mapping, planned rows, problems).

    Row status: 'ok', 'warning' (imported, some value ignored), 'duplicate' or
    'error' (both skipped). Each row keeps its validated values under 'cleaned'.
    """
    options = options or {}
    cols = _clean_mapping(mapping, len(table["headers"]))
    use_columns = cols["last_name"] is not None and cols["first_name"] is not None
    use_full = not use_columns and cols["full_name"] is not None
    if not use_columns and not use_full:
        return cols, [], [messages.IMPORT_NO_NAME_COLUMNS]

    default_establishment = collapse_spaces(options.get("establishment"))
    default_class = collapse_spaces(options.get("class_name"))
    first_name_first = bool(options.get("first_name_first"))
    overrides = _clean_overrides(overrides)

    def get(cells: list, field: str) -> Any:
        index = cols[field]
        return cells[index] if index is not None and index < len(cells) else None

    rows: list[dict] = []
    seen: dict[tuple[str, str], int] = {}
    for line, cells in table["rows"]:
        notes: list[str] = []
        if use_columns:
            last, first = cell_text(get(cells, "last_name")), cell_text(get(cells, "first_name"))
        else:
            last, first = split_full_name(cell_text(get(cells, "full_name")), first_name_first)
        last, first = smart_title(last), smart_title(first)

        sex, ok = parse_sex(get(cells, "sex"))
        if not ok:
            notes.append(messages.import_row_sex_ignored(cell_text(get(cells, "sex"))))
        birth_date, ok = parse_birth_date(get(cells, "birth_date"))
        if not ok:
            notes.append(messages.import_row_birth_ignored(cell_text(get(cells, "birth_date"))))
        age, ok = parse_age(get(cells, "age"))
        if not ok:
            notes.append(messages.import_row_age_ignored(cell_text(get(cells, "age"))))
        vma, ok = parse_vma(get(cells, "vma"))
        if not ok:
            notes.append(messages.import_row_vma_ignored(cell_text(get(cells, "vma"))))

        if line in overrides:
            class_name = overrides[line]
        else:
            class_name = cell_text(get(cells, "class_name")) or default_class
        establishment = cell_text(get(cells, "establishment")) or default_establishment

        raw = {
            "last_name": last, "first_name": first, "sex": sex, "class_name": class_name,
            "establishment": establishment, "birth_date": birth_date, "age": age,
            "vma": vma, "bib": None,
        }
        row = {"line": line, "status": "ok", "messages": notes, "values": raw, "cleaned": None}
        rows.append(row)

        if not last or not first:
            row["status"] = "error"
            row["messages"] = [messages.IMPORT_ROW_MISSING_NAME]
            continue
        try:
            cleaned = participants.clean_participant(raw)
        except ValueError as exc:  # e.g. a name longer than the limit
            row["status"] = "error"
            row["messages"] = [str(exc)]
            continue

        row["values"] = {**raw, "last_name": cleaned["last_name"], "first_name": cleaned["first_name"],
                         "class_name": cleaned["class"] or "", "establishment": cleaned["establishment"] or ""}
        key = (cleaned["last_name_key"], cleaned["first_name_key"])
        if key in existing:
            row["status"], row["messages"] = "duplicate", [messages.IMPORT_ROW_DUPLICATE_EXISTING]
        elif key in seen:
            row["status"], row["messages"] = "duplicate", [messages.import_row_duplicate_file(seen[key])]
        else:
            seen[key] = line
            row["cleaned"] = cleaned
            row["status"] = "warning" if notes else "ok"
    return cols, rows, []


def _existing_keys(event_file: str) -> set[tuple[str, str]]:
    with db.connect(existing_event_path(event_file)) as conn:
        return participants.identity_keys(conn)


def preview(event_file: str, table: dict, mapping: dict | None, options: dict | None,
            overrides: dict | None) -> dict:
    """What an import would do with the current mapping and options (nothing is written)."""
    cols, rows, problems = _plan(table, mapping, options, overrides, _existing_keys(event_file))
    count = lambda *statuses: sum(1 for r in rows if r["status"] in statuses)  # noqa: E731
    return {
        "mapping": cols,
        "problems": problems,
        "summary": {
            "total": len(rows),
            "importable": count("ok", "warning"),
            "warnings": count("warning"),
            "duplicates": count("duplicate"),
            "errors": count("error"),
        },
        "rows": [{k: r[k] for k in ("line", "status", "messages", "values")} for r in rows],
    }


def commit(event_file: str, table: dict, mapping: dict | None, options: dict | None,
           overrides: dict | None) -> dict:
    """Insert every importable row in one transaction (all or nothing)."""
    with db.connect(existing_event_path(event_file)) as conn:
        existing = participants.identity_keys(conn)
        _, rows, problems = _plan(table, mapping, options, overrides, existing)
        if problems:
            raise ValueError(problems[0])
        importable = [r for r in rows if r["status"] in ("ok", "warning")]
        if not importable:
            raise ValueError(messages.IMPORT_NOTHING)
        for row in importable:
            participants.insert_participant(conn, row["cleaned"])
    return {
        "imported": len(importable),
        "duplicates": sum(1 for r in rows if r["status"] == "duplicate"),
        "errors": sum(1 for r in rows if r["status"] == "error"),
    }


# --- template ------------------------------------------------------------------------------

TEMPLATE_HEADERS = ("Nom", "Prénom", "Sexe", "Classe", "Établissement", "Date de naissance", "Âge", "VMA")
TEMPLATE_NAME = "modele-participants-crossly.xlsx"
_TEMPLATE_NOTES = (
    "Comment remplir le modèle",
    "",
    "• Une ligne par participant, à partir de la ligne 2 de la feuille « Participants ».",
    "• Seuls le nom et le prénom sont obligatoires. Les autres colonnes peuvent rester vides.",
    "• Sexe : Garçon ou Fille (M, F, G sont aussi acceptés).",
    "• Classe : par exemple 6A, 5B, 43 (6e A ou 6 A sont reconnus et corrigés).",
    "• Établissement : s'il est le même pour toute la liste, laissez la colonne vide et "
    "saisissez-le au moment de l'import.",
    "• Date de naissance : jj/mm/aaaa. Si vous ne la connaissez pas, indiquez l'âge.",
    "• VMA : en km/h, entre 5 et 30 (virgule ou point). Laissez vide si elle est inconnue.",
    "• Les noms écrits tout en majuscules (export Pronote) sont corrigés à l'import.",
    "• Les doublons (même nom et même prénom) sont repérés et ignorés.",
)


def write_template(destination: str) -> str:
    """Write the downloadable .xlsx template and return its path."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    path = Path(destination)
    if path.suffix.lower() != ".xlsx":
        path = path.with_name(path.name + ".xlsx")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Participants"
    sheet.append(list(TEMPLATE_HEADERS))
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F6FEB")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for column, width in zip("ABCDEFGH", (22, 22, 12, 10, 28, 18, 8, 10)):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"

    sex = DataValidation(type="list", formula1='"Garçon,Fille"', allow_blank=True)
    sheet.add_data_validation(sex)
    sex.add("C2:C2000")
    for row in range(2, 2001):
        sheet.cell(row=row, column=6).number_format = "DD/MM/YYYY"
        sheet.cell(row=row, column=8).number_format = "0.0"

    notes = workbook.create_sheet("Mode d'emploi")
    for text in _TEMPLATE_NOTES:
        notes.append([text])
    notes["A1"].font = Font(bold=True, size=14)
    notes.column_dimensions["A"].width = 110

    workbook.save(path)
    return str(path)
